from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font

from storage import Storage


WEIGHT_HEADERS = ["Дата и время", "Вес, кг"]
WATER_HEADERS = ["Дата и время", "Вода, мл"]
PROFILE_HEADERS = ["Показатель", "Значение"]


def _get_or_create_sheet(wb: Workbook, title: str):
    return wb[title] if title in wb.sheetnames else wb.create_sheet(title)


def _reset_sheet(ws, headers: list[str]) -> None:
    ws.delete_rows(1, ws.max_row)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _parse_dt(value: str) -> datetime:
    return datetime.fromisoformat(value).replace(tzinfo=None)


def sync_excel(storage: Storage, user_id: int, excel_path: Path, water_norm_ml: int) -> None:
    if excel_path.exists():
        wb = load_workbook(excel_path)
    else:
        wb = Workbook()

    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]

    weight_ws = _get_or_create_sheet(wb, "Вес")
    water_ws = _get_or_create_sheet(wb, "Вода")
    profile_ws = _get_or_create_sheet(wb, "Профиль")

    _reset_sheet(weight_ws, WEIGHT_HEADERS)
    _reset_sheet(water_ws, WATER_HEADERS)
    _reset_sheet(profile_ws, PROFILE_HEADERS)

    weights = storage.all_weights(user_id)
    water_logs = storage.all_water_logs(user_id)

    for row in weights:
        weight_ws.append([_parse_dt(row["measured_at"]), row["weight_kg"]])

    for row in water_logs:
        water_ws.append([_parse_dt(row["logged_at"]), row["amount_ml"]])

    latest = weights[-1]["weight_kg"] if weights else None
    first = weights[0]["weight_kg"] if weights else None
    delta = round(latest - first, 2) if latest is not None and first is not None else None
    total_water = sum(int(row["amount_ml"]) for row in water_logs)
    user = storage.get_user(user_id)
    target = user["target_weight_kg"] if user and "target_weight_kg" in user.keys() else None
    avg_7_days = storage.average_weight_since(user_id, datetime.now(), 7)

    profile_ws.append(["Последний вес", latest])
    profile_ws.append(["Цель по весу", target])
    profile_ws.append(["До цели, кг", round(latest - target, 2) if latest is not None and target is not None else None])
    profile_ws.append(["Средний вес за 7 дней", avg_7_days])
    profile_ws.append(["Изменение веса от первой записи", delta])
    profile_ws.append(["Всего воды записано, мл", total_water])
    profile_ws.append(["Дневная норма воды, мл", water_norm_ml])

    for ws in (weight_ws, water_ws, profile_ws):
        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 18

    weight_ws._charts.clear()
    if len(weights) >= 1:
        chart = LineChart()
        chart.title = "График веса"
        chart.y_axis.title = "Вес, кг"
        chart.x_axis.title = "Дата"
        data = Reference(weight_ws, min_col=2, min_row=1, max_row=weight_ws.max_row)
        cats = Reference(weight_ws, min_col=1, min_row=2, max_row=weight_ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 9
        chart.width = 18
        weight_ws.add_chart(chart, "D2")

    wb.save(excel_path)
